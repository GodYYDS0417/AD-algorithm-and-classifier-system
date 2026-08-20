#!/usr/bin/env python3
"""Verify the generated docx content."""
from docx import Document

docx_path = r"C:\Users\28130\Desktop\lyy-6.12\doubao_final.docx"
doc = Document(docx_path)

print("Total paragraphs:", len(doc.paragraphs))
print("\n" + "=" * 60)
print("LAST 15 PARAGRAPHS (algorithm section):")
print("=" * 60)

for i, para in enumerate(doc.paragraphs[-15:]):
    idx = len(doc.paragraphs) - 15 + i
    if para.text.strip():
        preview = para.text[:100] + "..." if len(para.text) > 100 else para.text
        print(f"[{idx}] {preview}")
