#!/usr/bin/env python3
"""Read docx file content."""
from docx import Document

docx_path = r"C:\Users\28130\Desktop\lyy-6.12\实验方法部分.docx"
doc = Document(docx_path)

print("=" * 60)
print("DOCUMENT PARAGRAPHS:")
print("=" * 60)

for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        style = para.style.name if para.style else "Normal"
        print(f"[{i}] [{style}] {para.text}")

print("\n" + "=" * 60)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print("=" * 60)

if doc.tables:
    for t_idx, table in enumerate(doc.tables):
        print(f"\nTable {t_idx}:")
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            print(f"  | {' | '.join(row_text)} |")
